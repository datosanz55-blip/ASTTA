import os
import zipfile
from io import BytesIO

from openpyxl import Workbook, load_workbook
from PIL import Image

from catalog.importer import build_template, import_catalog
from catalog.models import Product
from tests.conftest import make_product


def filled_template(rows):
    wb = load_workbook(BytesIO(build_template()))
    ws = wb.worksheets[0]
    for row in rows:
        ws.append(row)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def photo(color="white", size=(2400, 1800), fmt="JPEG"):
    buffer = BytesIO()
    Image.new("RGB", size, color).save(buffer, fmt)
    return buffer.getvalue()


def archive(files):
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as zf:
        for name, data in files.items():
            zf.writestr(name, data)
    buffer.seek(0)
    return buffer


ROW = ["S-1", "Кроссовки беговые", "shoes", 120, 0.9, "5:100", "", "", "", "Лёгкие",
       "https://detail.1688.com/offer/1.html", ""]


def test_import_publishes_for_team(supplier, shoes):
    result = import_catalog(filled_template([
        ["S-1", "Кроссовки беговые", "Обувь", 120, "0,9", "5:100；20:90", "40-42", "чёрный, белый", "нет",
         "Сетка, EVA", "https://detail.1688.com/offer/1.html", "跑步鞋"],
        ["S-2", "Шлёпанцы", "鞋类", 30, 0.4, "", "", "", "", "", "", ""],
    ]), supplier, publish=True)
    assert result.ok, result.errors
    assert result.created == 2
    first = Product.objects.get(sku="S-1")
    assert first.status == Product.Status.PUBLISHED
    assert (first.title_ru, first.title_zh, first.description_ru) == ("Кроссовки беговые", "跑步鞋", "Сетка, EVA")
    assert first.source_url == "https://detail.1688.com/offer/1.html"
    assert (first.size_list, first.color_list) == (["40", "41", "42"], ["чёрный", "белый"])
    assert [(t.min_qty, int(t.price_cny)) for t in first.tiers.all()] == [(5, 100), (20, 90)]
    assert not Product.objects.get(sku="S-2").tiers.exists()  # без опта: только розница


def test_import_without_publish_goes_to_moderation(supplier, shoes):
    import_catalog(filled_template([ROW]), supplier)
    assert Product.objects.get(sku="S-1").status == Product.Status.MODERATION


def test_import_with_error_loads_nothing(supplier, shoes):
    result = import_catalog(filled_template([
        ROW,
        ["S-2", "Шлёпанцы", "unknown", "abc", 0.4, "10-25", "", "", "", "", "ftp://x", ""],
    ]), supplier)
    assert not result.ok
    assert any("категория" in e for e in result.errors)
    assert any("цена" in e for e in result.errors)
    assert any("6:198" in e for e in result.errors)
    assert any("ссылка" in e for e in result.errors)
    assert Product.objects.count() == 0


def test_reimport_updates_same_sku(supplier, shoes):
    import_catalog(filled_template([ROW]), supplier)
    result = import_catalog(filled_template([[*ROW[:3], 150, *ROW[4:]]]), supplier)
    assert (result.created, result.updated) == (0, 1)
    assert Product.objects.get(sku="S-1").price_cny == 150


def test_broken_file(supplier):
    result = import_catalog(BytesIO(b"not excel"), supplier)
    assert not result.ok


def test_photos_matched_by_sku_and_converted(supplier, shoes):
    photos = archive({
        "S-1_2.jpg": photo("red"),
        "S-1.jpg": photo("white"),
        "S-10/main.png": photo("blue", fmt="PNG"),
        "random.jpg": photo(),
        "notes.txt": b"hello",
        "__MACOSX/._S-1.jpg": b"junk",
    })
    result = import_catalog(filled_template([ROW, ["S-10", "Тапки", "shoes", 20, 0.3, "", "", "", "", "", "", ""]]),
                            supplier, photos=photos, publish=True)
    assert result.ok, result.errors
    assert result.photos == 3
    first = Product.objects.get(sku="S-1")
    images = list(first.images.all())
    assert [img.image.name.endswith(".webp") for img in images] == [True, True]
    # Порядок по имени: сначала S-1.jpg, потом S-1_2.jpg. Фото уменьшено до 1600 px.
    with Image.open(images[0].image.path) as main:
        assert main.size == (1600, 1200)
        assert main.getpixel((10, 10))[:3] == (255, 255, 255)
    assert Product.objects.get(sku="S-10").images.count() == 1  # папка S-10, а не артикул S-1
    assert any("random.jpg" in w for w in result.warnings)
    assert any("notes.txt" in w for w in result.warnings)


def test_photos_only_for_existing_products_replace_old(supplier, shoes, django_capture_on_commit_callbacks):
    product = make_product(supplier, shoes, "s-1", 100, 1, [(5, 80)])
    import_catalog(None, supplier, photos=archive({"S-1.jpg": photo("red"), "S-1-2.jpg": photo()}))
    assert product.images.count() == 2
    old = product.images.first().image.path
    with django_capture_on_commit_callbacks(execute=True):
        result = import_catalog(None, supplier, photos=archive({"s-1.jpg": photo("blue")}))
    assert result.ok and result.photos == 1
    assert product.images.count() == 1
    assert not os.path.exists(old)  # старый файл удалён после записи в базу


def test_broken_photo_loads_nothing(supplier, shoes):
    result = import_catalog(filled_template([ROW]), supplier, photos=archive({"S-1.jpg": b"not a photo"}))
    assert not result.ok
    assert any("S-1.jpg" in e for e in result.errors)
    assert Product.objects.count() == 0


def test_nothing_to_import(supplier):
    assert not import_catalog(None, supplier).ok
    assert not import_catalog(None, supplier, photos=BytesIO(b"not zip")).ok


def test_old_template_is_rejected(supplier, shoes):
    wb = Workbook()
    wb.active.append(["Артикул", "Название", "Категория", "Цена у поставщика, ¥", "Вес в упаковке, кг",
                      "Опт: от шт.:цена ¥; … (можно пусто)", "Крупногабарит (да/нет)"])
    wb.active.append(["S-1", "Тапки", "shoes", 10, 0.3, "", "нет"])
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    result = import_catalog(buffer, supplier)
    assert not result.ok and "свежий шаблон" in result.errors[0]
    assert not Product.objects.exists()
