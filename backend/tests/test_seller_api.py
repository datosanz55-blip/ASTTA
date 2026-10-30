import json
import zipfile
from decimal import Decimal
from io import BytesIO

import pytest
from django.contrib.auth.models import Group
from django.core.cache import cache
from django.test import Client
from openpyxl import load_workbook
from PIL import Image

from accounts.models import User
from accounts.roles import setup_roles
from catalog.importer import build_template
from catalog.models import Product, SupplierPost
from orders.models import Payout, SupplierOrder
from orders.services import CartLine, create_order, create_payouts, mark_received, supplier_action
from tests.conftest import make_product

API = "/api/seller"


def as_json(client, method, path, data=None, **extra):
    return getattr(client, method)(f"{API}{path}", data=json.dumps(data) if data is not None else None,
                                   content_type="application/json", **extra)


@pytest.fixture(autouse=True)
def _clear_throttle():
    cache.clear()


@pytest.fixture
def setup(supplier, other_supplier, shoes, pricing_setup):
    setup_roles()
    mine = make_product(supplier, shoes, "mine", 100, 1, [(5, 80), (20, 70)], title_ru="Мой товар", title_zh="我的鞋")
    theirs = make_product(other_supplier, shoes, "theirs", 200, 1, [(5, 150)], title_ru="Чужой товар")
    user = User.objects.create_user("jj_supplier", password="pw-123456!", role="supplier", supplier=supplier,
                                    language="zh-hans")
    user.groups.add(Group.objects.get(name="Поставщики"))
    return {"mine": mine, "theirs": theirs, "user": user, "supplier": supplier, "other": other_supplier}


@pytest.fixture
def client(setup):
    client = Client()
    client.force_login(setup["user"])
    return client


def test_page_is_served_in_live_mode():
    resp = Client().get("/seller/")
    html = resp.content.decode()
    assert resp.status_code == 200
    assert 'data-api="/api/seller"' in html
    assert 'src="assets/' not in html  # пути переписаны на static
    assert "csrftoken" in resp.cookies


def test_login_requires_csrf_and_valid_password(setup):
    client = Client(enforce_csrf_checks=True)
    client.get("/seller/")
    token = client.cookies["csrftoken"].value
    bad = as_json(client, "post", "/login", {"username": "jj_supplier", "password": "wrong"}, HTTP_X_CSRFTOKEN=token)
    assert bad.status_code == 401
    ok = as_json(client, "post", "/login", {"username": "jj_supplier", "password": "pw-123456!"},
                 HTTP_X_CSRFTOKEN=token)
    assert ok.status_code == 200
    assert ok.json()["supplier"]["code"] == "jj"
    # Запись без CSRF-токена отклоняется. После входа Django выдаёт новый токен, кабинет читает его из cookie.
    assert as_json(client, "put", "/me/language", {"language": "ru"}).status_code == 403
    assert as_json(client, "put", "/me/language", {"language": "ru"}, HTTP_X_CSRFTOKEN=token).status_code == 403
    fresh = client.cookies["csrftoken"].value
    assert as_json(client, "put", "/me/language", {"language": "ru"}, HTTP_X_CSRFTOKEN=fresh).status_code == 200


def test_login_is_throttled(setup):
    client = Client()
    codes = [as_json(client, "post", "/login", {"username": "x", "password": "y"}).status_code for _ in range(12)]
    assert codes[-1] == 429


def test_buyer_cannot_open_cabinet(setup):
    buyer = User.objects.create_user("buyer", password="pw-123456!")
    client = Client()
    client.force_login(buyer)
    assert client.get(f"{API}/me").status_code == 403
    assert client.get(f"{API}/products").status_code == 403
    assert Client().get(f"{API}/products").status_code == 401


def test_supplier_sees_only_own_products(client, setup):
    data = client.get(f"{API}/products").json()
    assert [p["zh"] for p in data] == ["我的鞋"]
    other = setup["theirs"]
    payload = {"zh": "改", "cat": "shoes", "price": "1", "kg": "1", "tiers": [[1, "1"]]}
    assert as_json(client, "put", f"/products/{other.pk}", payload).status_code == 404


def test_create_and_update_product_goes_to_moderation(client, setup):
    payload = {"zh": "新款拖鞋", "cat": "shoes", "price": "30", "kg": "0.4", "tiers": [[20, "16"], [100, "13"]],
               "desc": "轻便"}
    created = as_json(client, "post", "/products", payload).json()
    assert created["status"] == "moderation"
    assert created["ru"] == ""  # русского пока нет, модератор переведёт
    product = Product.objects.get(pk=created["id"])
    assert product.supplier == setup["supplier"]
    assert product.title_ru == "新款拖鞋"

    published = setup["mine"]
    updated = as_json(client, "put", f"/products/{published.pk}", {**payload, "zh": "我的鞋 新"}).json()
    assert updated["status"] == "moderation"
    assert [t[0] for t in updated["tiers"]] == [20, 100]


def test_product_validation(client):
    bad = {"zh": " ", "cat": "nope", "price": "0", "kg": "0", "tiers": []}
    resp = as_json(client, "post", "/products", bad)
    assert resp.status_code == 422
    for key in ("zh", "cat", "price", "kg", "tiers"):
        assert key in resp.json()["detail"]


def test_image_upload(client, setup):
    buffer = BytesIO()
    Image.new("RGB", (20, 20), "white").save(buffer, "PNG")
    buffer.seek(0)
    buffer.name = "shoe.png"
    resp = client.post(f"{API}/products/{setup['mine'].pk}/image", {"file": buffer})
    assert resp.status_code == 200
    assert resp.json()["img"].endswith(".webp")  # фото хранится уменьшенным в WebP
    not_image = BytesIO(b"not an image")
    not_image.name = "x.png"
    assert client.post(f"{API}/products/{setup['mine'].pk}/image", {"file": not_image}).status_code == 400


def test_orders_flow_without_buyer_contacts(client, setup):
    order = create_order([CartLine("mine", 1, "42"), CartLine("theirs", 1)], contact="@secret_buyer")
    data = client.get(f"{API}/orders").json()
    assert len(data) == 1  # только своя часть заказа
    so = data[0]
    assert so["num"] == order.number
    assert [i["zh"] for i in so["items"]] == ["我的鞋"]
    assert "secret_buyer" not in json.dumps(data)

    assert client.post(f"{API}/orders/{so['id']}/ship", data="{}", content_type="application/json").status_code == 422
    assert as_json(client, "post", f"/orders/{so['id']}/ship", {"tracking": "YT1"}).status_code == 409
    assert client.post(f"{API}/orders/{so['id']}/confirm").json()["status"] == "ship"
    assert as_json(client, "post", f"/orders/{so['id']}/ship", {"tracking": " "}).status_code == 409
    shipped = as_json(client, "post", f"/orders/{so['id']}/ship", {"tracking": "YT1"}).json()
    assert (shipped["status"], shipped["tracking"]) == ("shipped", "YT1")
    assert order.events.filter(note__contains="YT1").exists()

    theirs = SupplierOrder.objects.get(order=order, supplier=setup["other"])
    assert client.post(f"{API}/orders/{theirs.pk}/confirm").status_code == 404


def test_payouts_from_received_orders(client, setup):
    order = create_order([CartLine("mine", 2)], contact="@b")
    so = order.supplier_orders.get()
    assert client.get(f"{API}/overview").json()["pending_cny"] == "0"
    supplier_action(so, "confirm")
    mark_received(so)
    assert Decimal(client.get(f"{API}/overview").json()["pending_cny"]) == Decimal("200")  # 100 ¥ × 2, без сбора
    payouts = create_payouts()
    assert [p.amount_cny for p in payouts] == [Decimal("200")]
    data = client.get(f"{API}/payouts").json()
    assert Decimal(data["pending_cny"]) == 0
    assert data["items"][0]["orders"] == 1
    Payout.objects.update(status=Payout.Status.PAID)
    assert Decimal(client.get(f"{API}/payouts").json()["paid_total_cny"]) == Decimal("200")


def test_import_from_cabinet(client):
    wb = load_workbook(BytesIO(build_template()))
    wb.worksheets[0].append(["S-9", "Шлёпанцы", "shoes", 30, 0.4, "10:25", "", "", "", "拖鞋"])
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    buffer.name = "catalog.xlsx"
    # Поставщик не может выложить сразу на сайт, даже если попросит.
    result = client.post(f"{API}/import", {"file": buffer, "publish": "true"}).json()
    assert (result["ok"], result["created"]) == (True, 1)
    assert Product.objects.get(sku="S-9").status == Product.Status.MODERATION
    broken = BytesIO(b"x")
    broken.name = "bad.xlsx"
    assert client.post(f"{API}/import", {"file": broken}).json()["ok"] is False
    assert client.get(f"{API}/template")["Content-Type"].startswith("application/vnd.openxmlformats")


def test_shop_settings_cannot_rename_for_buyers(client, setup):
    resp = as_json(client, "put", "/shop", {"name_zh": "晋江鞋厂", "wechat": "jj_new", "payout_details": "Alipay 138"})
    data = resp.json()
    assert (data["name_zh"], data["wechat"], data["payout_details"]) == ("晋江鞋厂", "jj_new", "Alipay 138")
    assert data["name"] == "TT · Цзиньцзян"


def test_factory_and_posts(client, setup):
    assert client.get(f"{API}/factory").json()["per_day"] == 3
    saved = as_json(client, "put", "/factory", {"new_items": False, "weekly": True, "restock": True, "per_day": 5})
    assert saved.json() == {"new_items": False, "weekly": True, "restock": True, "per_day": 5}
    too_many = as_json(client, "put", "/factory", {"new_items": True, "weekly": True, "restock": True, "per_day": 50})
    assert too_many.status_code == 422
    post = client.post(f"{API}/posts", {"product": setup["mine"].pk, "kind": "new", "text": "新到"}).json()
    assert post["status"] == "moderation"
    assert SupplierPost.objects.get().supplier == setup["supplier"]
    foreign = client.post(f"{API}/posts", {"product": setup["theirs"].pk, "kind": "new", "text": "x"})
    assert foreign.status_code == 404
    assert Client().get("/api/feed").json()["total"] == 0  # пост поставщика ждёт модерации


def test_moderator_can_open_any_cabinet(setup):
    moderator = User.objects.create_user("mod", password="pw-123456!", is_staff=True)
    moderator.groups.add(Group.objects.get(name="Модераторы"))
    client = Client()
    client.force_login(moderator)
    me = client.get(f"{API}/me").json()
    assert me["can_act"] is True
    assert {s["code"] for s in me["suppliers"]} == {"jj", "gz"}
    assert client.get(f"{API}/products").status_code == 400  # нужно выбрать поставщика
    gz = client.get(f"{API}/products", HTTP_X_ACT_AS="gz").json()
    assert [p["slug"] for p in gz] == ["theirs"]


def test_blocked_supplier_is_locked_out(client, setup):
    setup["supplier"].is_active = False
    setup["supplier"].save()
    assert client.get(f"{API}/products").status_code == 403


@pytest.fixture
def team(setup):
    moderator = User.objects.create_user("mod", password="pw-123456!", is_staff=True)
    moderator.groups.add(Group.objects.get(name="Модераторы"))
    client = Client()
    client.force_login(moderator)
    return client


def png(color="white"):
    buffer = BytesIO()
    Image.new("RGB", (40, 30), color).save(buffer, "PNG")
    buffer.seek(0)
    buffer.name = f"{color}.png"
    return buffer


def test_team_creates_product_in_russian_and_publishes(team, setup):
    payload = {"ru": "Тапки домашние", "sku": "T-1", "cat": "shoes", "price": "12", "kg": "0.3", "desc": "Мягкие",
               "src": "https://detail.1688.com/offer/2.html"}
    resp = as_json(team, "post", "/products", payload, HTTP_X_ACT_AS="jj")
    assert resp.status_code == 200, resp.content
    data = resp.json()
    assert (data["status"], data["ru"], data["zh"], data["desc"]) == ("published", "Тапки домашние", "", "Мягкие")
    product = Product.objects.get(pk=data["id"])
    assert (product.sku, product.source_url, product.tiers.count()) == ("T-1", payload["src"], 0)
    assert as_json(team, "post", "/products", {**payload, "ru": ""}, HTTP_X_ACT_AS="jj").status_code == 422
    duplicate = as_json(team, "post", "/products", payload, HTTP_X_ACT_AS="jj")
    assert duplicate.status_code == 422 and "sku" in duplicate.json()["detail"]
    photo = team.post(f"{API}/products/{product.pk}/image", {"file": png()}, HTTP_X_ACT_AS="jj")
    assert photo.json()["status"] == "published"  # фото от команды не отправляет товар на модерацию


def test_team_import_publishes_with_photos(team, setup):
    wb = load_workbook(BytesIO(build_template()))
    wb.worksheets[0].append(["S-9", "Шлёпанцы", "shoes", 30, 0.4, "", "", "", "", ""])
    sheet = BytesIO()
    wb.save(sheet)
    sheet.seek(0)
    sheet.name = "catalog.xlsx"
    photos = BytesIO()
    with zipfile.ZipFile(photos, "w") as zf:
        zf.writestr("S-9.png", png().getvalue())
        zf.writestr("other.png", png().getvalue())
    photos.seek(0)
    photos.name = "photos.zip"
    result = team.post(f"{API}/import", {"file": sheet, "photos": photos, "publish": "true"}, HTTP_X_ACT_AS="jj").json()
    assert (result["ok"], result["created"], result["photos"]) == (True, 1, 1), result
    assert any("other.png" in w for w in result["warnings"])
    product = Product.objects.get(sku="S-9")
    assert product.status == Product.Status.PUBLISHED and product.supplier == setup["supplier"]


def test_team_post_with_photos_goes_to_feed(team, setup):
    resp = team.post(f"{API}/posts", {"product": setup["mine"].pk, "kind": "new", "text": "Новинка: мягкие тапки",
                                      "images": [png("red"), png("blue")]}, HTTP_X_ACT_AS="jj")
    assert resp.status_code == 200, resp.content
    post = resp.json()
    assert post["status"] == "published" and len(post["images"]) == 2
    saved = SupplierPost.objects.get()
    assert (saved.text_ru, saved.text) == ("Новинка: мягкие тапки", "")

    feed = Client().get("/api/feed").json()
    assert feed["total"] == 1
    item = feed["items"][0]
    assert item["text"] == "Новинка: мягкие тапки" and len(item["images"]) == 2
    assert item["product"]["slug"] == "mine" and item["product"]["retail_rub"]
    assert Client().get("/api/feed?supplier=gz").json()["total"] == 0

    # Снятый с продажи товар из поста пропадает, сам пост остаётся.
    Product.objects.filter(pk=setup["mine"].pk).update(status=Product.Status.MODERATION)
    assert Client().get("/api/feed").json()["items"][0]["product"] is None

    photos_only = team.post(f"{API}/posts", {"text": "", "images": [png()]}, HTTP_X_ACT_AS="jj")
    assert photos_only.status_code == 200  # пост без товара и без текста, только фото
    assert team.post(f"{API}/posts", {"text": " "}, HTTP_X_ACT_AS="jj").status_code == 422
    assert team.delete(f"{API}/posts/{post['id']}", HTTP_X_ACT_AS="jj").status_code == 200
    assert Client().get("/api/feed").json()["total"] == 1
